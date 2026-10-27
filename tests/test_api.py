"""Tests de l'API. Lancer :  pytest -q"""
import cv2
import numpy as np
import pytest
from fastapi.testclient import TestClient

from api.main import app, SCENARIOS
from kromamed.simulate import simulate_photo

client = TestClient(app)


def test_sante():
    r = client.get("/api/health")
    assert r.status_code == 200 and r.json()["status"] == "ok"


def test_page_d_accueil():
    r = client.get("/")
    assert r.status_code == 200 and "KromaMed" in r.text and "/static/app.js" in r.text


def test_application_installable_manifest_et_service_worker():
    m = client.get("/manifest.webmanifest")
    assert m.status_code == 200
    data = m.json()
    assert data["display"] == "standalone" and data["theme_color"] == "#F2B134"
    for icon in data["icons"]:                       # chaque icône annoncée existe vraiment
        assert client.get(icon["src"]).status_code == 200
    sw = client.get("/sw.js")
    assert sw.status_code == 200 and "javascript" in sw.headers["content-type"]
    assert "/api/" in sw.text                        # le service worker exclut l'API du cache


def test_ecrans_et_fichiers_de_l_application():
    html = client.get("/").text
    for ecran in ("s-home", "s-test", "s-camera", "s-result", "s-history", "s-guide", "s-profile"):
        assert f'id="{ecran}"' in html
    assert "Profil" in html and "Réglages" not in html
    for fichier in ("/static/app.js", "/static/app.css", "/static/maquette_mobile.png",
                    "/static/fonts/plus-jakarta-sans-latin-700-normal.woff2"):
        assert client.get(fichier).status_code == 200


def test_liste_des_scenarios():
    ids = {s["id"] for s in client.get("/api/scenarios").json()}
    assert ids == set(SCENARIOS)


def test_demo_proteinurie_detectee():
    d = client.get("/api/demo/proteinurie").json()
    assert d["status"] == "ok"
    assert d["summary"]["level"] == "alert"
    assert len(d["pads"]) == 10
    pro = next(p for p in d["pads"] if p["param"] == "proteines")
    assert pro["level"] == "++" and pro["label"] == "Protéines" and pro["read_time_s"] == 60
    assert d["images"]["input"].startswith("data:image/jpeg;base64,")


@pytest.mark.parametrize("name", ["flou", "ombre_forte", "repere_cache"])
def test_demos_degradees_sont_refusees(name):
    d = client.get(f"/api/demo/{name}").json()
    assert d["status"] == "rejected" and d["pads"] == []
    assert d["summary"]["level"] == "rejected" and d["summary"]["details"]


def test_rapport_pdf():
    d = client.get("/api/demo/infection").json()
    d.pop("images")
    r = client.post("/api/report", json={"patient_id": "TEST_001", "date": "2026-10-19 12:00", "result": d})
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF") and len(r.content) > 5000


def test_rapport_refuse_si_analyse_refusee():
    d = client.get("/api/demo/flou").json()
    d.pop("images")
    r = client.post("/api/report", json={"patient_id": "X", "result": d})
    assert r.status_code == 422


def test_export_excel_colonnes_du_modele():
    from io import BytesIO
    from openpyxl import load_workbook
    rows = [{"date": "2026-10-19 12:00", "patient_id": "TEST_001",
             "results": {"ph": "6,0", "densite": "1,015", "leucocytes": "négatif"}}]
    r = client.post("/api/history/xlsx", json={"rows": rows})
    assert r.status_code == 200
    ws = load_workbook(BytesIO(r.content)).active
    header = [c.value for c in ws[1]]
    assert header[:4] == ["Date_Test", "Patient_ID", "Marque_Bandelette", "1_Leucocytes"] and len(header) == 13
    line = [c.value for c in ws[2]]
    assert line[1] == "TEST_001" and line[3] == "Négatif" and line[7] == "6.0" and line[9] == "1.015"


def test_scenario_inconnu():
    assert client.get("/api/demo/nimporte").status_code == 404


def test_envoi_d_une_photo():
    img = simulate_photo({"proteines": "+"}, seed=5)
    ok, buf = cv2.imencode(".jpg", img)
    r = client.post("/api/analyze", files={"file": ("photo.jpg", buf.tobytes(), "image/jpeg")})
    assert r.status_code == 200
    d = r.json()
    assert d["status"] == "ok"
    assert next(p for p in d["pads"] if p["param"] == "proteines")["level"] == "+"


def test_fichier_illisible():
    r = client.post("/api/analyze", files={"file": ("x.txt", b"pas une image", "text/plain")})
    assert r.status_code == 400


def test_photo_sans_carte_refusee_proprement():
    ok, buf = cv2.imencode(".png", np.full((400, 600, 3), 120, np.uint8))
    d = client.post("/api/analyze", files={"file": ("v.png", buf.tobytes(), "image/png")}).json()
    assert d["status"] == "rejected"
