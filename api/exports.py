"""Export Excel de l'historique, au format de base_donnees_tests.xlsx. Généré en mémoire."""
from __future__ import annotations
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from kromamed import config as C

COLUMN_NAMES = {
    "leucocytes": "1_Leucocytes", "nitrites": "2_Nitrites", "urobilinogene": "3_Urobilinogene",
    "proteines": "4_Proteines", "ph": "5_pH", "sang": "6_Sang", "densite": "7_Densite_Specifique",
    "cetones": "8_Cetones", "bilirubine": "9_Bilirubine", "glucose": "10_Glucose",
}


def _cell(param: str, level: str | None) -> str:
    if not level:
        return "Indéterminé"
    text = level.capitalize() if level in ("négatif", "normal", "traces", "positif") else level
    return text.replace(",", ".") if param in ("ph", "densite") else text


def build_history_xlsx(rows: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    header = ["Date_Test", "Patient_ID", "Marque_Bandelette"] + [COLUMN_NAMES[p] for p in C.PARAMS]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B1F3A")
        cell.alignment = Alignment(horizontal="center")
    for r in rows:
        res = r.get("results", {})
        ws.append([r.get("date", ""), r.get("patient_id", "") or "ANONYME", C.BRAND]
                  + [_cell(p, res.get(p)) for p in C.PARAMS])
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(12, min(26, max(len(str(c.value or "")) for c in col) + 2))
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
